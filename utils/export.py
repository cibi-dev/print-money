# utils/export.py – Exportación de reportes a Excel y PDF
# Usa openpyxl para Excel (import perezoso) y ReportLab para PDF

from datetime import datetime
import db_manager
from utils.i18n import T


def exportar_excel(ruta_salida: str, año: int):
    """
    Exporta un reporte anual completo a Excel con múltiples hojas.

    Hojas:
    - Pedidos: todos los pedidos del año
    - Gastos: todos los gastos del año
    - Resumen mensual: KPIs por mes
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    fecha_inicio = f"{año}-01-01"
    fecha_fin    = f"{año}-12-31"

    pedidos = db_manager.obtener_pedidos(filtro_fecha_desde=fecha_inicio,
                                         filtro_fecha_hasta=fecha_fin)
    gastos  = db_manager.obtener_gastos(fecha_desde=fecha_inicio, fecha_hasta=fecha_fin)

    resumen_meses = db_manager.obtener_resumen_mensual_anual(año)
    meses = T("months_full")

    wb = Workbook()
    ws_resumen = wb.active
    ws_resumen.title = T("reports_sheet_summary")
    ws_resumen.append([
        T("reports_col_month"), T("reports_col_income_currency"),
        T("reports_col_expenses_currency"), T("reports_col_profit_currency"),
        T("reports_col_margin_pct"), T("reports_col_orders"),
    ])
    for mes in range(1, 13):
        stats = resumen_meses[mes]
        ws_resumen.append([
            meses[mes - 1],
            stats["ingresos"],
            stats["gastos"],
            stats["ganancia_neta"],
            round(stats["margen"], 2),
            stats["num_pedidos"],
        ])

    ws_pedidos = wb.create_sheet(T("reports_sheet_orders"))
    cols_ped = ["id", "fecha", "cliente_nombre", "descripcion", "filamento_nombre",
                "gramos_usados", "horas_impresion", "costo_material", "costo_tiempo",
                "costo_envio", "precio_venta", "ganancia", "margen_porcentaje",
                "fecha_entrega", "estado"]
    ws_pedidos.append(cols_ped)
    for pedido in pedidos:
        ws_pedidos.append([pedido.get(col) for col in cols_ped])

    ws_gastos = wb.create_sheet(T("reports_sheet_expenses"))
    cols_gas = ["id", "fecha", "categoria_nombre", "descripcion", "monto",
                "recurrente", "frecuencia"]
    ws_gastos.append(cols_gas)
    for gasto in gastos:
        ws_gastos.append([gasto.get(col) for col in cols_gas])

    for ws in (ws_resumen, ws_pedidos, ws_gastos):
        for celda in ws[1]:
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor="0F0F0F")
            celda.alignment = Alignment(horizontal="center")
        ws.freeze_panes = "A2"
        for col in ws.columns:
            max_len = max((len(str(c.value)) for c in col if c.value is not None),
                          default=10)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(
                max_len + 4, 40)

    wb.save(ruta_salida)


def exportar_pdf_reporte(ruta_salida: str, año: int):
    """
    Exporta un reporte anual resumido en PDF.
    Incluye tabla de KPIs mensuales y totales del año.
    """
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer,
                                    Table, TableStyle, HRFlowable)
    from reportlab.lib.colors import HexColor
    from reportlab.lib.enums import TA_CENTER

    COLOR_NEGRO  = HexColor("#0F0F0F")
    COLOR_BLANCO = HexColor("#FFFFFF")
    COLOR_VERDE  = HexColor("#00FF9D")
    COLOR_GRIS   = HexColor("#F5F5F5")

    doc = SimpleDocTemplate(ruta_salida, pagesize=A4,
                            topMargin=2*cm, bottomMargin=2*cm,
                            leftMargin=2*cm, rightMargin=2*cm)
    story = []
    styles = getSampleStyleSheet()
    W = A4[0] - 4*cm

    # Título
    story.append(Paragraph(
        T("reports_pdf_title", year=año),
        ParagraphStyle("t", fontSize=20, fontName="Helvetica-Bold",
                       textColor=COLOR_NEGRO, alignment=TA_CENTER, spaceAfter=6)
    ))
    story.append(Paragraph(
        T("reports_generated", date=datetime.now().strftime("%d/%m/%Y %H:%M")),
        ParagraphStyle("s", fontSize=9, textColor=HexColor("#888888"),
                       alignment=TA_CENTER, spaceAfter=4)
    ))
    story.append(Spacer(1, 0.4*cm))
    story.append(HRFlowable(width=W, thickness=2, color=COLOR_NEGRO))
    story.append(Spacer(1, 0.5*cm))

    # Tabla mensual
    resumen_meses = db_manager.obtener_resumen_mensual_anual(año)
    meses = T("months_full")
    headers = [T("reports_col_month"), T("reports_col_income"),
               T("reports_col_expenses"), T("reports_col_profit"),
               T("reports_col_margin"), T("reports_col_orders")]
    rows = [headers]
    total_ingresos = total_gastos = total_ganancia = 0
    for mes in range(1, 13):
        stats = resumen_meses[mes]
        total_ingresos += stats["ingresos"]
        total_gastos   += stats["gastos"]
        total_ganancia += stats["ganancia_neta"]
        rows.append([
            meses[mes - 1],
            f"${stats['ingresos']:,.2f}",
            f"${stats['gastos']:,.2f}",
            f"${stats['ganancia_neta']:,.2f}",
            f"{stats['margen']:.1f}%",
            str(stats["num_pedidos"]),
        ])

    # Fila de totales
    margen_total = (total_ganancia / total_ingresos * 100) if total_ingresos > 0 else 0
    rows.append([
        T("reports_total_row"), f"${total_ingresos:,.2f}", f"${total_gastos:,.2f}",
        f"${total_ganancia:,.2f}", f"{margen_total:.1f}%", ""
    ])

    tabla = Table(rows, colWidths=[W/6]*6)
    tabla.setStyle(TableStyle([
        ("BACKGROUND",    (0, 0), (-1, 0), COLOR_NEGRO),
        ("TEXTCOLOR",     (0, 0), (-1, 0), COLOR_BLANCO),
        ("FONTNAME",      (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE",      (0, 0), (-1, -1), 8),
        ("ALIGN",         (0, 0), (-1, -1), "CENTER"),
        ("TOPPADDING",    (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("ROWBACKGROUNDS",(0, 1), (-1, -2), [COLOR_GRIS, COLOR_BLANCO]),
        ("FONTNAME",      (0, -1), (-1, -1), "Helvetica-Bold"),
        ("BACKGROUND",    (0, -1), (-1, -1), HexColor("#E8FFF8")),
        ("GRID",          (0, 0), (-1, -1), 0.4, HexColor("#DDDDDD")),
    ]))
    story.append(tabla)
    story.append(Spacer(1, 1*cm))
    story.append(Paragraph(T("reports_pdf_footer"),
                           ParagraphStyle("pie", fontSize=10, fontName="Helvetica-Bold",
                                         textColor=COLOR_NEGRO, alignment=TA_CENTER)))

    doc.build(story)
